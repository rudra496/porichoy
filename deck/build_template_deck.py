# -*- coding: utf-8 -*-
"""Fill the official NIC 3.0 pitch-deck template with Porichoy content.
Keeps template design/logos/headings; replaces instruction & sample text;
embeds verified figures; deletes the instructions page; saves as
Porichoy_Pitch_Deck.pptx (their required naming: startup name)."""
from pptx import Presentation
from pptx.util import Emu, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
from pptx.enum.shapes import MSO_SHAPE
from PIL import Image
import copy

INK = RGBColor(0x1E, 0x1B, 0x4B)
GREY = RGBColor(0x4B, 0x48, 0x77)
PINK = RGBColor(0xE1, 0x1D, 0x74)
INDIGO = RGBColor(0x4F, 0x46, 0xE5)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
GREEN = RGBColor(0x15, 0x80, 0x3D)

SRC = 'template.pptx'
OUT = 'Porichoy_Pitch_Deck.pptx'
prs = Presentation(SRC)


def set_lines(shape, lines, align=None):
    """Replace a textbox content with styled paragraphs.
    lines: list of (text, size, color, bold, italic) or (text, size, color)."""
    tf = shape.text_frame
    tf.clear()
    first = True
    for item in lines:
        text, size, color = item[0], item[1], item[2]
        bold = item[3] if len(item) > 3 else False
        italic = item[4] if len(item) > 4 else False
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        run = p.add_run()
        run.text = text
        run.font.size = Pt(size)
        run.font.color.rgb = color
        run.font.bold = bold
        run.font.italic = italic
        run.font.name = 'Calibri'
        if align is not None:
            p.alignment = align


def find_box(slide, contains, name_contains=None):
    for sh in slide.shapes:
        if sh.has_text_frame:
            t = sh.text_frame.text
            if contains.lower() in t.lower() and (name_contains is None or name_contains in sh.name):
                return sh
    return None


def remove_logo_groups(slide):
    """Delete the template's pink 'YOUR LOGO (If Any)' placeholder GROUPS."""
    for sh in list(slide.shapes):
        if sh.shape_type == 6 and sh.left == 10560676 and sh.top == 5679584:
            sh._element.getparent().remove(sh._element)


def logo_replace(slide):
    for sh in list(slide.shapes):
        if sh.has_text_frame and sh.text_frame.text.strip() == 'YOUR LOGO':
            sh._element.getparent().remove(sh._element)


def add_img(slide, path, left, top, width=None, height=None):
    return slide.shapes.add_picture(path, Emu(int(left)), Emu(int(top)),
                                    Emu(int(width)) if width else None,
                                    Emu(int(height)) if height else None)


def img_size(path):
    with Image.open(path) as im:
        return im.size


def chip(slide, x, y, w, h, big, small, color):
    sh = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Emu(x), Emu(y), Emu(w), Emu(h))
    sh.fill.solid()
    sh.fill.fore_color.rgb = INK
    sh.line.fill.background()
    sh.shadow.inherit = False
    tf = sh.text_frame
    tf.clear()
    p = tf.paragraphs[0]
    r = p.add_run(); r.text = big
    r.font.size = Pt(17); r.font.bold = True; r.font.color.rgb = PINK; r.font.name = 'Calibri'
    p2 = tf.add_paragraph()
    r2 = p2.add_run(); r2.text = small
    r2.font.size = Pt(10); r2.font.color.rgb = WHITE; r2.font.name = 'Calibri'
    return sh


IN = '_work_imgs'
slides = prs.slides

# ---------- SLIDE 1 (instructions) -> delete later; record index 0 ----------
# ---------- SLIDE 2: Title ----------
s = slides[1]
remove_logo_groups(s)
title = find_box(s, 'TITLE OF DECK GOES HERE')
set_lines(title, [('Porichoy (\u09aa\u09b0\u09bf\u099a\u09df)', 40, INK, True)])
sub = find_box(s, 'Instruction: It can be the name')
set_lines(sub, [('DPP-in-a-Box \u2014 every Bangladeshi garment gets a verifiable identity before the EU\u2019s 2027 Digital Product Passport wall.', 16, GREY, False, True)])
tb = s.shapes.add_textbox(Emu(1524000), Emu(3300000), Emu(9144000), Emu(400000))
set_lines(tb, [('rudra496.github.io/porichoy  \u00b7  Sylhet, Bangladesh  \u00b7  Rudra Sarker, SUST IPE', 13, INDIGO, True)])

# ---------- SLIDE 3: THE PROBLEM ----------
s = slides[2]
logo_replace(s)
remove_logo_groups(s)
sample = find_box(s, 'Sample')
if sample is not None:
    sample._element.getparent().remove(sample._element)
box = find_box(s, "It\u2019s observed in cutting production flow")
set_lines(box, [
    ('EU customs will auto-check Digital Product Passports on imported garments: ESPR in force 18 Jul 2024 \u00b7 textiles in the first Working Plan (Apr 2025) \u00b7 textile delegated act expected ~2027. No passport data \u2192 no EU order.', 14, INK, True),
    ('The data already exists in every factory \u2014 trapped in mixed Bangla/English Excels, ERP exports and paper trim cards. Global traceability platforms are brand-side and enterprise-priced; nobody captures at the Bangladeshi factory floor.', 13, GREY),
    ('Most affected: small & mid-size exporters without compliance teams \u2014 and, through them, the \u22484 million garment workers (BGMEA; majority women) whose livelihoods ride on Europe, the industry\u2019s biggest market.', 13, GREY),
])
w, h = img_size(f'{IN}/chart_timeline.png')
iw = 9800000
ih = int(iw * h / w)
add_img(s, f'{IN}/chart_timeline.png', 1750000, 4150000, width=iw, height=ih)

# ---------- SLIDE 4: YOUR SOLUTION ----------
s = slides[3]
logo_replace(s)
remove_logo_groups(s)
box = find_box(s, 'In here, you need to tell us about your startup')
set_lines(box, [
    ('Porichoy turns a factory\u2019s existing records into buyer-ready EU Digital Product Passports in four steps \u2014 all live today, offline-tolerant, Bangla + English:', 14, INK, True),
    ('1 \u00b7 INGEST \u2014 drop the factory Excel/CSV; every sheet parsed on-device. Nothing leaves the factory (local-first).', 12.5, GREY),
    ('2 \u00b7 MAP \u2014 the engine reads chaotic Bangla/English headers with confidence scores; a trained steward confirms; the factory\u2019s vocabulary is learned.', 12.5, GREY),
    ('3 \u00b7 SCORE \u2014 DPP Readiness 0\u2013100 per production order (weights = ESPR Art. 8 categories) with the exact gap list.', 12.5, GREY),
    ('4 \u00b7 PASSPORT \u2014 one click \u2192 a QR whose link embeds the passport; scanning opens the buyer view on any phone. No backend.', 12.5, GREY),
    ('19/19 automated engine tests \u00b7 SHA-256 provenance chain \u00b7 stewards fix gaps in-place and the score updates instantly.', 12, PINK, True),
])
for i, nm in enumerate(['porichoy_dash.png', 'porichoy_po.png', 'porichoy_evidence.png']):
    w, h = img_size(f'{IN}/{nm}')
    iw = 3400000
    ih = int(iw * h / w)
    add_img(s, f'{IN}/{nm}', 1700000 + i * 3450000, 4550000, width=iw, height=ih)
tb = s.shapes.add_textbox(Emu(456353), Emu(6660000), Emu(11085000), Emu(190000))
set_lines(tb, [('Live production app \u2014 screenshots fetched 2026-09-13 \u00b7 rudra496.github.io/porichoy', 10, GREY, False, True)])

# ---------- SLIDE 5: MARKET SIZE & OPPORTUNITY ----------
s = slides[4]
logo_replace(s)
remove_logo_groups(s)
box = find_box(s, 'Here, you need to tell us about the market size')
set_lines(box, [
    ('A dated compliance wall creates a forced-adoption market:', 14, INK, True),
    ('\u2248 4 million RMG workers (BGMEA) \u00b7 world\u2019s #2 ready-made garment exporter \u00b7 Europe is the biggest market (NIC 3.0 call) \u00b7 thousands of export factories in the BGMEA member base \u2014 every one shipping to the EU must be DPP-ready inside the 2027\u201330 window.', 13, GREY),
    ('Illustrative revenue math on Porichoy\u2019s own pricing (BDT 2,000\u20138,000/factory/month):', 13, INK, True),
    ('100 pilot-scale factories \u2248 BDT 60 lakh/year \u00b7 1,000 factories \u2248 BDT 6 crore/year \u2014 before brand-funded supplier-onboarding contracts, where buyers pay to bring their supplier base onto passports before the deadline.', 13, GREY),
])
chip(s, 456353, 4350000, 2650000, 1050000, '\u22484 million', 'RMG workers, majority women (BGMEA)', PINK)
chip(s, 3306353, 4350000, 2650000, 1050000, '#2 exporter', 'world RMG \u2014 Europe the biggest market', INDIGO)
chip(s, 6156353, 4350000, 2650000, 1050000, '2027\u201330', 'EU DPP phase-in; customs auto-checks', PINK)
chip(s, 9006353, 4350000, 2650000, 1050000, '10\u2192100+', 'factories: year-1 pilot \u2192 year-3 scale', GREEN)
tb = s.shapes.add_textbox(Emu(456353), Emu(5550000), Emu(10700000), Emu(950000))
set_lines(tb, [
    ('Timing is the wedge: buyers already ask, budgets activate before the deadline, and no Bangladesh-built tool serves the factory floor today \u2014 the verified competitors are brand-side only.', 12, GREY),
])

# ---------- SLIDE 6: BUSINESS & REVENUE MODEL ----------
s = slides[5]
logo_replace(s)
remove_logo_groups(s)
box = find_box(s, 'Here, you need to describe how you are making money')
set_lines(box, [
    ('Three simple revenue lines:', 14, INK, True),
    ('1. Factory subscriptions \u2014 BDT 2,000\u20138,000/month by size (SME-priced, ~10\u00d7 below enterprise platforms).', 12.5, GREY),
    ('2. Onboarding & data-rescue \u2014 one-time fee: steward training + first ingest + mapping historic files.', 12.5, GREY),
    ('3. Brand-funded supplier onboarding \u2014 buyers pay to make their supplier base passport-ready before 2027\u201330.', 12.5, GREY),
    ('Unit economics: onboarding \u2248 2 person-days/factory \u00b7 software gross margin >80% \u00b7 replication is a template, not a project.', 12.5, PINK, True),
])
w, h = img_size(f'{IN}/chart_grant.png')
ih = 2500000
iw = int(ih * w / h)
add_img(s, f'{IN}/chart_grant.png', 6900000, 3950000, width=iw, height=ih)
tb = s.shapes.add_textbox(Emu(456353), Emu(3950000), Emu(5900000), Emu(1400000))
set_lines(tb, [
    ('Pollination grant (BDT 650,000) funds the 10-factory pilot:', 12.5, INK, True),
    ('BDT 220k pilot + steward training \u00b7 180k OCR for paper trim cards \u00b7 150k mapping corpus (Bangla + ERP formats) \u00b7 100k legal & compliance.', 12, GREY),
])

# ---------- SLIDE 7: YOUR IMPACT ----------
s = slides[6]
logo_replace(s)
remove_logo_groups(s)
box = find_box(s, 'Here, you have to explain')
set_lines(box, [
    ('Competitiveness: EU-order eligibility is protected (customs auto-checks), compliance prep per PO drops from days of file-hunting to under an hour, and readiness scores make factories the buyer-preferred choice.', 12.5, GREY),
    ('Women: two women operators per factory are trained & certified as DPP Data Stewards \u2014 formal, higher-skilled roles (20 in year 1 \u2192 200 by year 3, targets) \u2014 and keeping factories order-eligible protects a majority-women workforce (\u22484M workers, BGMEA).', 12.5, GREY),
    ('Circularity: DPP fields make recycled content, certificates and end-of-life routes verifiable \u2014 turning circularity claims into machine-checkable data that flows to buyers and recyclers.', 12.5, GREY),
])
w, h = img_size(f'{IN}/chart_women.png')
iw = 4400000
ih = int(iw * h / w)
add_img(s, f'{IN}/chart_women.png', 1700000, 3900000, width=iw, height=ih)
w2, h2 = img_size(f'{IN}/photo_women.jpg')
ih2 = 2500000
iw2 = int(ih2 * w2 / h2)
add_img(s, f'{IN}/photo_women.jpg', 6350000, 3950000, width=iw2, height=ih2)
tb = s.shapes.add_textbox(Emu(456353), Emu(6600000), Emu(11085000), Emu(300000))
set_lines(tb, [('Left: women-steward ramp (roadmap targets). Right: RMG factory floor \u2014 photo CC BY 2.0, Wikimedia Commons.', 10, GREY, False, True)])

# ---------- SLIDE 8: YOUR TEAM/YOURSELF ----------
s = slides[7]
logo_replace(s)
remove_logo_groups(s)
box = find_box(s, 'Do you have the credibility to pull this forward')
set_lines(box, [
    ('Rudra Sarker \u2014 Founder & full-time engineer. Final-year B.Sc. Industrial & Production Engineering, SUST. Factory-process education + production AI engineering: exactly this product\u2019s stack.', 14, INK, True),
    ('Ships at production grade, solo: AdalatAI \u2014 AI-assisted court platform, 30-type trilingual classifier (F1 0.963) \u00b7 JolSetu \u2014 arsenic-testing PWA, live map of 6,593 real wells \u00b7 RippleUp \u2014 Android + Windows app, release v5.3.3 \u00b7 a self-trained LLM serving system on a VPS.', 12.5, GREY),
    ('Published researcher: Journal of Ethnopharmacology (Q1, 2026), 3rd author.', 12.5, GREY),
    ('No big-name advisers \u2014 instead: 19/19 passing engine tests, public code (github.com/rudra496/porichoy), a live product, and honest staging: launched Sep 2026, pre-revenue until the pilot.', 12.5, PINK, True),
])
w, h = img_size(f'{IN}/rudra_cv_photo.png')
add_img(s, f'{IN}/rudra_cv_photo.png', 1750000, 4750000, width=1800000, height=1800000)
for i, nm in enumerate(['adalatai.png', 'jolsetu.png']):
    w, h = img_size(f'{IN}/{nm}')
    iw = 2700000
    ih = int(iw * h / w)
    add_img(s, f'{IN}/{nm}', 6300000 + i * 2900000, 4900000, width=iw, height=ih)
tb = s.shapes.add_textbox(Emu(556797), Emu(6600000), Emu(11085000), Emu(230000))
set_lines(tb, [('Right: live productions \u2014 AdalatAI (court AI) and JolSetu (safe-water), both built and shipped by the founder.', 10, GREY, False, True)])

# ---------- SLIDE 9: THANK YOU ----------
s = slides[8]
logo_replace(s)
remove_logo_groups(s)
logo_replace(s)
remove_logo_groups(s)
box = find_box(s, 'Donald Trump')
set_lines(box, [
    ('Name:  Rudra Sarker \u2014 Founder, Porichoy', 16, INK, True),
    ('Phone:  +8801988223165', 14, GREY),
    ('Email:  rudrasarker125@gmail.com', 14, GREY),
    ('Live product:  rudra496.github.io/porichoy  \u00b7  Code:  github.com/rudra496/porichoy', 14, INDIGO, True),
])

# ---------- delete instructions slide ----------
xml_slides = prs.slides._sldIdLst
sld = list(xml_slides)[0]
xml_slides.remove(sld)

prs.save(OUT)
print('saved', OUT, 'slides:', len(prs.slides.__iter__.__self__._sldIdLst))
